import argparse
import json
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.session import SessionLocal
from app.models.farm import Farm

DEFAULT_SHEET = "Food_Donors"
REQUIRED_COLUMNS = (
    "donorname",
    "streetaddress",
    "city",
    "state",
)


@dataclass(frozen=True)
class FarmImportRow:
    name: str
    address: str | None
    city: str | None
    region: str | None


def normalize_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value)
    normalized = normalized.casefold()
    return re.sub(r"\s+", " ", normalized).strip()


def normalize_header(value: object) -> str:
    if value is None:
        return ""

    normalized = unicodedata.normalize("NFKC", str(value))
    normalized = normalized.casefold()
    return re.sub(r"[^a-z0-9]+", "", normalized)


def clean_cell(value: object) -> str | None:
    if value is None:
        return None

    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    return cleaned or None


def load_farm_rows(
    file_path: Path,
    sheet_name: str = DEFAULT_SHEET,
) -> list[FarmImportRow]:
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(
                f"worksheet {sheet_name!r} was not found; available sheets: {available}"
            )

        worksheet = workbook[sheet_name]
        row_iterator = worksheet.iter_rows(values_only=True)
        header_row = next(row_iterator, None)

        if header_row is None:
            raise ValueError(f"worksheet {sheet_name!r} is empty")

        column_indices: dict[str, int] = {}

        for index, header in enumerate(header_row):
            normalized_header = normalize_header(header)

            if normalized_header:
                column_indices.setdefault(normalized_header, index)

        missing_columns = set(REQUIRED_COLUMNS) - set(column_indices)

        if missing_columns:
            missing = ", ".join(sorted(missing_columns))
            raise ValueError(f"donor worksheet is missing columns: {missing}")

        rows: list[FarmImportRow] = []
        seen_names: dict[str, int] = {}

        for row_number, values in enumerate(row_iterator, start=2):
            name = clean_cell(values[column_indices["donorname"]])

            if name is None:
                continue

            normalized_name = normalize_name(name)

            if normalized_name in seen_names:
                first_row = seen_names[normalized_name]
                raise ValueError(
                    f"row {row_number}: duplicate donor name; first seen on row {first_row}"
                )

            seen_names[normalized_name] = row_number

            rows.append(
                FarmImportRow(
                    name=name,
                    address=clean_cell(values[column_indices["streetaddress"]]),
                    city=clean_cell(values[column_indices["city"]]),
                    region=clean_cell(values[column_indices["state"]]),
                )
            )

        return rows
    finally:
        workbook.close()


def synchronize_farms(
    db: Session,
    rows: list[FarmImportRow],
    commit: bool,
) -> dict[str, object]:
    farms_by_name: dict[str, Farm] = {}

    for farm in db.scalars(select(Farm)).all():
        normalized_name = normalize_name(farm.name)

        if normalized_name in farms_by_name:
            raise ValueError(f"duplicate farm already exists in database: {farm.name}")

        farms_by_name[normalized_name] = farm

    farms_created = 0
    farms_updated = 0
    farms_unchanged = 0

    for row in rows:
        normalized_name = normalize_name(row.name)
        farm = farms_by_name.get(normalized_name)

        if farm is None:
            farm = Farm(
                name=row.name,
                address=row.address,
                city=row.city,
                region=row.region,
                is_active=True,
            )
            db.add(farm)
            farms_by_name[normalized_name] = farm
            farms_created += 1
            continue

        changed = False

        if farm.name != row.name:
            farm.name = row.name
            changed = True

        for field_name in ("address", "city", "region"):
            source_value = getattr(row, field_name)

            if source_value is not None and getattr(farm, field_name) != source_value:
                setattr(farm, field_name, source_value)
                changed = True

        if changed:
            farms_updated += 1
        else:
            farms_unchanged += 1

    db.flush()

    summary: dict[str, object] = {
        "source_rows": len(rows),
        "farms_created": farms_created,
        "farms_updated": farms_updated,
        "farms_unchanged": farms_unchanged,
        "committed": commit,
    }

    if commit:
        db.commit()
    else:
        db.rollback()

    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--file",
        required=True,
        type=Path,
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
    )
    parser.add_argument(
        "--commit",
        action="store_true",
    )
    arguments = parser.parse_args()

    rows = load_farm_rows(
        file_path=arguments.file,
        sheet_name=arguments.sheet,
    )

    with SessionLocal() as db:
        try:
            summary = synchronize_farms(
                db=db,
                rows=rows,
                commit=arguments.commit,
            )
        except Exception:
            db.rollback()
            raise

    summary["file"] = str(arguments.file)
    summary["sheet"] = arguments.sheet

    print(
        json.dumps(
            summary,
            indent=2,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()